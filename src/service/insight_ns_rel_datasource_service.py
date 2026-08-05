from concurrent.futures import ThreadPoolExecutor, TimeoutError, as_completed
from contextlib import contextmanager
from pathlib import Path
from threading import BoundedSemaphore
from uuid import uuid4
from typing import Any

import pandas as pd
from werkzeug.datastructures import FileStorage

from api import supos_kernel_api
from config.config import Config
from config.database import release_clean_session_connection
from sqlalchemy.orm import Session

from dto import DataSourceSchema, PropertySchema
from model import InsightDatasource, InsightNsConversation, InsightNsRelDatasource, InsightNsUnsSelection
from service.datasource_description_runner import (
    DESCRIPTION_SOURCE_DEFAULT,
    DESCRIPTION_SOURCE_USER,
    build_description_meta,
    get_datasource_description_runner,
    schema_guard_hash,
)
from utils import (
    TABLE_ENGINE_RELATIONAL,
    TABLE_ENGINE_TIME_SERIES,
    dump_json,
    normalize_datasource_type,
    normalize_ts_query_source,
    safe_json_loads,
    logger,
)

SHARED_UNS_NAMESPACE_ID = 0
DEFAULT_CONVERSATION_ID = 0
BIND_SOURCE_USER_SELECTED = 'user_selected'
BIND_SOURCE_IMPORT_SELECTED = 'import_selected'
BIND_SOURCE_SYSTEM_DEFAULT = 'system_default'
UNS_TIME_SERIES_DATA_TYPE = 1
UNS_RELATIONAL_TABLE_DATA_TYPE = 2
UNS_CALCULATED_TABLE_DATA_TYPE = 3
UNS_REFERENCE_TABLE_DATA_TYPE = 7
UNS_PROTOCOL_TABLE_DATA_TYPE = 9
UPLOADED_DATASOURCE_DESCRIPTION_PENDING = '见解自动生成中，请稍后查看'
UNS_TABLE_DATA_TYPES = {
    UNS_TIME_SERIES_DATA_TYPE,
    UNS_RELATIONAL_TABLE_DATA_TYPE,
    UNS_CALCULATED_TABLE_DATA_TYPE,
    UNS_REFERENCE_TABLE_DATA_TYPE,
    UNS_PROTOCOL_TABLE_DATA_TYPE,
}
UNS_MAX_EXPANDED_FILES = Config.UNS_MAX_EXPANDED_FILES
UNS_MAX_EXPAND_DEPTH = Config.UNS_MAX_EXPAND_DEPTH
UNS_TREE_PAGE_SIZE = Config.UNS_TREE_PAGE_SIZE
UNS_DETAIL_WORKERS = Config.UNS_DETAIL_WORKERS
UNS_IMPORT_MAX_CONCURRENT = Config.UNS_IMPORT_MAX_CONCURRENT
UNS_IMPORT_SEMAPHORE = BoundedSemaphore(UNS_IMPORT_MAX_CONCURRENT)


@contextmanager
def _uns_import_slot():
    acquired = UNS_IMPORT_SEMAPHORE.acquire(blocking=False)
    if not acquired:
        raise RuntimeError('UNS 节点导入操作繁忙，请稍后再试')
    try:
        yield
    finally:
        UNS_IMPORT_SEMAPHORE.release()


class InsightNsRelDatasourceService:
    """管理空间级数据源以及会话级绑定关系。"""

    def __init__(self, session: Session):
        self.session = session

    def find_by_conversation_id(self, insight_conversation_id: int) -> list[dict[str, Any]]:
        """查询某个会话已绑定的数据源。"""
        rows = self.session.query(InsightNsRelDatasource, InsightDatasource).filter(
            InsightNsRelDatasource.insight_conversation_id == insight_conversation_id,
            InsightNsRelDatasource.datasource_id == InsightDatasource.id,
            InsightNsRelDatasource.is_deleted == 0,
            InsightDatasource.is_deleted == 0,
        ).order_by(
            InsightNsRelDatasource.sort_no.asc(),
            InsightNsRelDatasource.id.asc(),
        ).all()
        return [self._to_dict(rel, datasource) for rel, datasource in rows]

    def find_by_namespace_id(
        self,
        insight_namespace_id: int,
        insight_conversation_id: int | None = None,
    ) -> list[dict[str, Any]]:
        """
        查询空间级数据源列表。

        如果传入会话 ID，返回结果会直接附带 `checked`，供前端勾选框展示。
        """
        bound_datasource_ids: set[int] = set()
        if insight_conversation_id:
            bound_rows = self.session.query(
                InsightNsRelDatasource.datasource_id,
                InsightNsRelDatasource.bind_source,
            ).filter(
                InsightNsRelDatasource.insight_conversation_id == insight_conversation_id,
                InsightNsRelDatasource.is_deleted == 0,
            ).all()
            bound_datasource_ids = {
                int(row[0])
                for row in bound_rows
                if (row[1] or BIND_SOURCE_USER_SELECTED) != BIND_SOURCE_SYSTEM_DEFAULT
            }

        rows = self.session.query(InsightDatasource).filter(
            InsightDatasource.is_deleted == 0,
            InsightDatasource.insight_namespace_id == insight_namespace_id,
        ).order_by(
            InsightDatasource.created_at.desc(),
            InsightDatasource.id.desc(),
        ).all()

        return [
            self._datasource_to_dict(
                row,
                checked=row.id in bound_datasource_ids,
                insight_conversation_id=insight_conversation_id or 0,
            )
            for row in rows
        ]

    def bind_existing_datasource(self, insight_conversation_id: int, datasource_id: int) -> dict[str, Any]:
        """把一条空间级数据源绑定到会话。"""
        conversation = self._get_conversation(insight_conversation_id)
        if conversation is None:
            return {"success": False, "message": "会话不存在"}

        datasource = self.session.query(InsightDatasource).filter(
            InsightDatasource.id == datasource_id,
            InsightDatasource.insight_namespace_id == conversation.insight_namespace_id,
            InsightDatasource.is_deleted == 0,
        ).first()
        if datasource is None:
            return {"success": False, "message": "数据源不存在"}

        existing = self.session.query(InsightNsRelDatasource).filter(
            InsightNsRelDatasource.insight_conversation_id == insight_conversation_id,
            InsightNsRelDatasource.datasource_id == datasource.id,
            InsightNsRelDatasource.is_deleted == 0,
        ).first()
        if existing is not None:
            return {"success": True, "message": "数据源已绑定到当前会话", "data": self._to_dict(existing, datasource)}

        relation = InsightNsRelDatasource(
            insight_namespace_id=conversation.insight_namespace_id,
            insight_conversation_id=insight_conversation_id,
            datasource_id=datasource.id,
            is_active=1,
            sort_no=self._next_sort_no(insight_conversation_id),
            bind_source=BIND_SOURCE_USER_SELECTED,
        )
        self.session.add(relation)
        self.session.commit()
        return {"success": True, "message": "数据源已绑定到当前会话", "data": self._to_dict(relation, datasource)}

    def update_namespace_datasource_description(
        self,
        insight_namespace_id: int,
        datasource_id: int,
        description: str,
    ) -> dict[str, Any]:
        """
        更新空间数据源描述。

        描述存放在 datasource_schema.description，因此这里只改这一项，
        避免编辑描述时覆盖其他 schema 元信息。
        """
        datasource = self.session.query(InsightDatasource).filter(
            InsightDatasource.id == datasource_id,
            InsightDatasource.insight_namespace_id.in_([SHARED_UNS_NAMESPACE_ID, insight_namespace_id]),
            InsightDatasource.is_deleted == 0,
        ).first()
        if datasource is None:
            return {"success": False, "message": "数据源不存在"}

        schema_payload = safe_json_loads(datasource.datasource_schema, {})
        if not isinstance(schema_payload, dict):
            schema_payload = {}
        schema_payload['description'] = str(description or '').strip()
        config_payload = safe_json_loads(datasource.datasource_config_json, {})
        if not isinstance(config_payload, dict):
            config_payload = {}
        current_meta = config_payload.get('description_meta') if isinstance(config_payload.get('description_meta'), dict) else {}
        config_payload['description_meta'] = build_description_meta(
            DESCRIPTION_SOURCE_USER,
            int(current_meta.get('version') or 0) + 1,
        )

        datasource.datasource_schema = dump_json(schema_payload)
        datasource.datasource_config_json = dump_json(config_payload)
        self.session.commit()
        self.session.refresh(datasource)
        return {
            "success": True,
            "message": "数据源描述已更新",
            "data": self._datasource_to_dict(datasource),
        }

    def get_uploaded_file_datasource_content(
        self,
        insight_namespace_id: int,
        datasource_id: int,
        page_no: int = 1,
        page_size: int = 50,
    ) -> dict[str, Any]:
        datasource = self.session.query(InsightDatasource).filter(
            InsightDatasource.id == datasource_id,
            InsightDatasource.insight_namespace_id == insight_namespace_id,
            InsightDatasource.datasource_type == 'local_file',
            InsightDatasource.is_deleted == 0,
        ).first()
        if datasource is None:
            return {"success": False, "message": "数据源不存在或不是外部上传文件"}

        config_payload = safe_json_loads(datasource.datasource_config_json, {})
        if not isinstance(config_payload, dict):
            config_payload = {}
        file_path_text = str(config_payload.get('file_path') or '').strip()
        if not file_path_text:
            return {"success": False, "message": "数据源文件路径不存在"}

        file_path = Path(file_path_text)
        if not file_path.exists() or not file_path.is_file():
            return {"success": False, "message": "数据源文件不存在或已被移动"}

        normalized_page_no = max(int(page_no or 1), 1)
        max_page_size = max(int(getattr(Config, 'DATASOURCE_FILE_CONTENT_MAX_PAGE_SIZE', 200) or 200), 1)
        normalized_page_size = min(max(int(page_size or 50), 1), max_page_size)
        sheet_name = str(config_payload.get('sheet_name') or '').strip() or None

        try:
            page_data = self._read_file_content_page(
                file_path=file_path,
                sheet_name=sheet_name,
                page_no=normalized_page_no,
                page_size=normalized_page_size,
            )
        except ValueError as exc:
            return {"success": False, "message": str(exc)}

        schema_payload = safe_json_loads(datasource.datasource_schema, {})
        if not isinstance(schema_payload, dict):
            schema_payload = {}

        page_data.update({
            "datasource_id": datasource.id,
            "datasource_name": datasource.datasource_name,
            "datasource_description": str(schema_payload.get('description') or '').strip(),
            "file_name": str(config_payload.get('original_filename') or file_path.name),
            "sheet_name": sheet_name or '',
            "page_no": normalized_page_no,
            "page_size": normalized_page_size,
            "max_page_size": max_page_size,
        })
        return {"success": True, "message": "数据源文件内容读取成功", "data": page_data}

    def import_uns_nodes_to_namespace(
        self,
        insight_namespace_id: int,
        insight_conversation_id: int = 0,
        nodes: list[dict[str, Any]] | None = None,
        authorization: str = '',
        lake_rds_database_name: str = '',
        ts_query_source: str = '',
    ) -> dict[str, Any]:
        """
        将选中的 UNS 节点转换为共享 table 数据源，并绑定到当前会话。

        UNS 数据源本体只在虚拟共享空间落一份；当前会话通过关系表获得自己的
        绑定记录，分析链路仍然只按当前会话绑定查询。
        """
        if insight_namespace_id <= 0:
            return {"success": False, "message": "空间不存在"}
        conversation = self._get_conversation(insight_conversation_id)
        if conversation is None or conversation.insight_namespace_id != insight_namespace_id:
            return {"success": False, "message": "当前会话不存在或不属于该空间"}
        release_clean_session_connection(self.session)

        normalized_nodes = self._normalize_uns_nodes(nodes)
        normalized_ts_query_source = normalize_ts_query_source(ts_query_source)

        try:
            with _uns_import_slot():
                imported_rows: list[dict[str, Any]] = []
                failed_rows: list[dict[str, Any]] = []
                desired_selections: list[dict[str, Any]] = []
                desired_file_node_ids: list[str] = []
                desired_datasource_ids: list[int] = []

                if normalized_nodes:
                    if not authorization:
                        return {"success": False, "message": "缺少 SUPOS 认证信息，无法导入 UNS 节点"}

                    expand_result = self._expand_uns_file_nodes(normalized_nodes, authorization)
                    if expand_result["failed"]:
                        return {
                            "success": False,
                            "message": self._build_uns_failure_message(
                                expand_result["failed"],
                                "UNS \u8282\u70b9\u540c\u6b65\u5931\u8d25\uff0c\u8bf7\u7a0d\u540e\u91cd\u8bd5",
                            ),
                            "data": {
                                "imported": [],
                                "failed": expand_result["failed"],
                                "selections": self.list_uns_selections(insight_conversation_id),
                            },
                        }

                    expanded_file_nodes = expand_result["file_nodes"]
                    node_type_by_id = {
                        str(node.get("id") or "").strip(): node.get("data_type")
                        for node in expanded_file_nodes
                        if str(node.get("id") or "").strip()
                    }
                    lake_rds_required_types = UNS_TABLE_DATA_TYPES - {UNS_PROTOCOL_TABLE_DATA_TYPE}
                    if any(data_type in lake_rds_required_types for data_type in node_type_by_id.values()) \
                            and not lake_rds_database_name:
                        return {
                            "success": False,
                            "message": "当前用户上下文未初始化 LakeRDS 数据库名",
                        }
                    time_series_required_types = {UNS_TIME_SERIES_DATA_TYPE, UNS_CALCULATED_TABLE_DATA_TYPE}
                    if any(data_type in time_series_required_types for data_type in node_type_by_id.values()) \
                            and not normalized_ts_query_source:
                        return {
                            "success": False,
                            "message": "当前平台时序数据库类型尚未初始化，无法导入时序 UNS 节点",
                            "data": {
                                "imported": [],
                                "failed": [
                                    {
                                        "id": node_id,
                                        "message": "缺少平台时序数据库类型",
                                    }
                                    for node_id, data_type in node_type_by_id.items()
                                    if data_type in time_series_required_types
                                ],
                                "selections": self.list_uns_selections(insight_conversation_id),
                            },
                        }

                    desired_selections = self._canonicalize_uns_selections(expand_result["selections"])
                    desired_file_node_ids = self._normalize_uns_node_id_list([
                        node.get("id")
                        for selection in desired_selections
                        for node in [{"id": item} for item in (selection.get("expanded_uns_node_ids") or [])]
                    ])
                    if not desired_file_node_ids:
                        self._replace_uns_selections(
                            insight_namespace_id=insight_namespace_id,
                            insight_conversation_id=insight_conversation_id,
                            selections=[],
                        )
                        removed_datasource_ids = self._sync_uns_conversation_bindings(
                            insight_namespace_id=insight_namespace_id,
                            insight_conversation_id=insight_conversation_id,
                            desired_datasource_ids=[],
                        )
                        selections = self.list_uns_selections(insight_conversation_id)
                        self.session.commit()
                        return {
                            "success": True,
                            "message": "已清空 UNS 节点选择",
                            "data": {
                                "imported": [],
                                "failed": [],
                                "removed_datasource_ids": removed_datasource_ids,
                                "selections": selections,
                            },
                        }

                    existing_datasources = self._find_uns_datasources_by_node_ids(desired_file_node_ids)
                    sync_node_ids = [
                        node_id
                        for node_id in desired_file_node_ids
                    ]
                    if sync_node_ids:
                        release_clean_session_connection(self.session)
                        detail_results = self._fetch_uns_details_in_parallel(
                            sync_node_ids,
                            authorization,
                        )
                        failed_rows.extend([
                            {
                                "id": result.get("id", ""),
                                "message": result.get("message", "导入失败"),
                            }
                            for result in detail_results
                            if not result.get("success")
                        ])
                        if failed_rows:
                            return {
                                "success": False,
                                "message": self._build_uns_failure_message(
                                    failed_rows,
                                    "UNS 节点同步失败，请稍后重试",
                                ),
                                "data": {
                                    "imported": [],
                                    "failed": failed_rows,
                                    "selections": self.list_uns_selections(insight_conversation_id),
                                },
                            }

                        reference_detail_by_id, reference_failed_rows = self._fetch_reference_uns_details_for_import(
                            [result["detail"] for result in detail_results if result.get("success")],
                            authorization,
                        )
                        failed_rows.extend(reference_failed_rows)
                        if failed_rows:
                            return {
                                "success": False,
                                "message": self._build_uns_failure_message(
                                    failed_rows,
                                    "UNS 引用节点同步失败，请稍后重试",
                                ),
                                "data": {
                                    "imported": [],
                                    "failed": failed_rows,
                                    "selections": self.list_uns_selections(insight_conversation_id),
                                },
                            }

                        for result in detail_results:
                            node_id = str(result.get("id") or (result.get("detail") or {}).get("id") or "").strip()
                            try:
                                datasource = self._upsert_uns_table_datasource(
                                    detail=result["detail"],
                                    lake_rds_database_name=lake_rds_database_name,
                                    uns_data_type=node_type_by_id.get(node_id),
                                    ts_query_source=normalized_ts_query_source,
                                    reference_detail_by_id=reference_detail_by_id,
                                )
                            except ValueError as exc:
                                failed_rows.append({
                                    "id": node_id,
                                    "message": str(exc),
                                })
                                continue
                            self._bind_datasource(
                                insight_namespace_id=SHARED_UNS_NAMESPACE_ID,
                                insight_conversation_id=DEFAULT_CONVERSATION_ID,
                                datasource_id=datasource.id,
                                bind_source=BIND_SOURCE_SYSTEM_DEFAULT,
                            )
                            if node_id not in existing_datasources:
                                imported_rows.append(self._datasource_to_dict(datasource))

                    if failed_rows:
                        return {
                            "success": False,
                            "message": self._build_uns_failure_message(
                                failed_rows,
                                "UNS 节点同步失败，请稍后重试",
                            ),
                            "data": {
                                "imported": imported_rows,
                                "failed": failed_rows,
                                "selections": self.list_uns_selections(insight_conversation_id),
                            },
                        }

                    desired_datasource_ids = self._find_datasource_ids_by_uns_node_ids(desired_file_node_ids)

                self._replace_uns_selections(
                    insight_namespace_id=insight_namespace_id,
                    insight_conversation_id=insight_conversation_id,
                    selections=desired_selections,
                )
                removed_datasource_ids = self._sync_uns_conversation_bindings(
                    insight_namespace_id=insight_namespace_id,
                    insight_conversation_id=insight_conversation_id,
                    desired_datasource_ids=desired_datasource_ids,
                )
        except RuntimeError as exc:
            return {"success": False, "message": str(exc)}

        selections = self.list_uns_selections(insight_conversation_id)
        self.session.commit()
        message = "UNS 节点选择已同步"
        if imported_rows:
            message = f"UNS 节点选择已同步，新增导入 {len(imported_rows)} 个数据源"
        return {
            "success": True,
            "message": message,
            "data": {
                "imported": imported_rows,
                "failed": failed_rows,
                "removed_datasource_ids": removed_datasource_ids,
                "selections": selections,
            },
        }

    def materialize_semantic_uns_node_ids(
        self,
        *,
        uns_node_ids: list[str],
        authorization: str,
        lake_rds_database_name: str = '',
        ts_query_source: str = '',
    ) -> dict[str, int]:
        """Materialize semantic-center node IDs without changing user selections."""
        normalized_node_ids = list(dict.fromkeys(
            str(item or '').strip() for item in uns_node_ids if str(item or '').strip()
        ))
        if not normalized_node_ids:
            raise ValueError('semantic selection returned no UNS node IDs')
        if not authorization:
            raise ValueError('semantic UNS materialization requires authorization')

        rows = self.session.query(InsightDatasource).filter(
            InsightDatasource.datasource_type == 'table',
            InsightDatasource.is_deleted == 0,
            InsightDatasource.uns_node_id.in_(normalized_node_ids),
        ).all()
        datasource_by_node_id = {
            str(datasource.uns_node_id or '').strip(): datasource
            for datasource in rows
            if str(datasource.uns_node_id or '').strip()
        }
        missing_node_ids = [
            node_id
            for node_id in normalized_node_ids
            if node_id not in datasource_by_node_id
        ]
        if missing_node_ids:
            release_clean_session_connection(self.session)
            detail_results = self._fetch_uns_details_in_parallel(
                missing_node_ids,
                authorization,
            )
            failed = [item for item in detail_results if not item.get('success')]
            if failed:
                raise ValueError(self._build_uns_failure_message(
                    [
                        {
                            'id': item.get('id', ''),
                            'message': item.get('message', 'UNS detail lookup failed'),
                        }
                        for item in failed
                    ],
                    'semantic UNS materialization failed',
                ))
            details: list[dict[str, Any]] = []
            for item in detail_results:
                requested_node_id = str(item.get('id') or '').strip()
                detail = dict(item['detail'])
                returned_node_id = str(detail.get('id') or '').strip()
                if returned_node_id and returned_node_id != requested_node_id:
                    raise ValueError(
                        'UNS instance response node ID mismatch: '
                        f'{requested_node_id} -> {returned_node_id}'
                    )
                detail['id'] = requested_node_id
                details.append(detail)
            data_types = {
                str(detail.get('id') or '').strip(): self._normalize_uns_data_type(detail)
                for detail in details
            }
            lake_required_types = UNS_TABLE_DATA_TYPES - {UNS_PROTOCOL_TABLE_DATA_TYPE}
            if any(item in lake_required_types for item in data_types.values()) \
                    and not lake_rds_database_name:
                raise ValueError('LakeRDS database name is unavailable')
            time_series_types = {UNS_TIME_SERIES_DATA_TYPE, UNS_CALCULATED_TABLE_DATA_TYPE}
            normalized_ts_query_source = normalize_ts_query_source(ts_query_source)
            if any(item in time_series_types for item in data_types.values()) \
                    and not normalized_ts_query_source:
                raise ValueError('time-series query source is unavailable')

            reference_detail_by_id, reference_failures = (
                self._fetch_reference_uns_details_for_import(details, authorization)
            )
            if reference_failures:
                raise ValueError(self._build_uns_failure_message(
                    reference_failures,
                    'semantic UNS reference materialization failed',
                ))
            for detail in details:
                datasource = self._upsert_uns_table_datasource(
                    detail=detail,
                    lake_rds_database_name=lake_rds_database_name,
                    uns_data_type=self._normalize_uns_data_type(detail),
                    ts_query_source=normalized_ts_query_source,
                    reference_detail_by_id=reference_detail_by_id,
                )
                self._bind_datasource(
                    insight_namespace_id=SHARED_UNS_NAMESPACE_ID,
                    insight_conversation_id=DEFAULT_CONVERSATION_ID,
                    datasource_id=datasource.id,
                    bind_source=BIND_SOURCE_SYSTEM_DEFAULT,
                )
                datasource_by_node_id[str(detail.get('id') or '').strip()] = datasource

        missing = [
            node_id
            for node_id in normalized_node_ids
            if node_id not in datasource_by_node_id
        ]
        if missing:
            raise ValueError(f'UNS node IDs could not be materialized: {missing}')
        self.session.commit()
        return {
            node_id: int(datasource_by_node_id[node_id].id)
            for node_id in normalized_node_ids
        }

    def _build_uns_failure_message(
        self,
        failed_rows: list[dict[str, Any]],
        default_message: str,
    ) -> str:
        if not failed_rows:
            return default_message

        first_message = str((failed_rows[0] or {}).get("message") or "").strip()
        if not first_message:
            return default_message
        if len(failed_rows) == 1:
            return first_message
        return (
            f"{first_message}\uff1b"
            f"\u53e6\u6709 {len(failed_rows) - 1} \u4e2a\u8282\u70b9\u5904\u7406\u5931\u8d25"
        )


    def list_uns_selections(self, insight_conversation_id: int) -> list[dict[str, Any]]:
        rows = self.session.query(InsightNsUnsSelection).filter(
            InsightNsUnsSelection.insight_conversation_id == insight_conversation_id,
            InsightNsUnsSelection.is_deleted == 0,
        ).order_by(
            InsightNsUnsSelection.id.asc(),
        ).all()
        selected_rows: list[dict[str, Any]] = []
        selected_node_ids: set[str] = set()
        derived_partial_ids: set[str] = set()

        for row in rows:
            payload = row.to_dict()
            payload["selection_state"] = "selected"
            payload["derived"] = False
            selected_rows.append(payload)
            selected_node_ids.add(str(payload.get("uns_node_id") or "").strip())

        partial_rows: list[dict[str, Any]] = []
        for payload in selected_rows:
            for ancestor_id in self._normalize_uns_node_id_list(payload.get("uns_node_path") or []):
                if ancestor_id in selected_node_ids or ancestor_id in derived_partial_ids:
                    continue
                partial_rows.append({
                    "id": f"partial:{ancestor_id}",
                    "insight_namespace_id": payload.get("insight_namespace_id"),
                    "insight_conversation_id": payload.get("insight_conversation_id"),
                    "uns_node_id": ancestor_id,
                    "uns_node_name": "",
                    "uns_node_path": [],
                    "is_folder": True,
                    "expanded_uns_node_ids": [],
                    "selection_state": "partial",
                    "derived": True,
                    "is_deleted": 0,
                    "created_at": payload.get("created_at"),
                    "updated_at": payload.get("updated_at"),
                })
                derived_partial_ids.add(ancestor_id)
        return selected_rows + partial_rows

    def delete_namespace_datasource(self, insight_namespace_id: int, datasource_id: int) -> dict[str, Any]:
        """删除空间级数据源，并自动解除该数据源的所有会话引用。"""
        datasource = self.session.query(InsightDatasource).filter(
            InsightDatasource.id == datasource_id,
            InsightDatasource.insight_namespace_id == insight_namespace_id,
            InsightDatasource.is_deleted == 0,
        ).first()
        if datasource is None:
            return {"success": False, "message": "数据源不存在"}

        self.session.query(InsightNsRelDatasource).filter(
            InsightNsRelDatasource.insight_namespace_id == insight_namespace_id,
            InsightNsRelDatasource.datasource_id == datasource_id,
            InsightNsRelDatasource.is_deleted == 0,
        ).update(
            {InsightNsRelDatasource.is_deleted: 1},
            synchronize_session=False,
        )

        datasource.is_deleted = 1
        self.session.commit()
        return {"success": True, "message": "数据源删除成功"}

    def remove_datasource(self, insight_conversation_id: int, datasource_id: int) -> dict[str, Any]:
        """从会话解绑数据源，不删除空间级数据源本体。"""
        row = self.session.query(InsightNsRelDatasource, InsightDatasource).filter(
            InsightNsRelDatasource.insight_conversation_id == insight_conversation_id,
            InsightNsRelDatasource.datasource_id == InsightDatasource.id,
            InsightDatasource.id == datasource_id,
            InsightNsRelDatasource.is_deleted == 0,
            InsightDatasource.is_deleted == 0,
        ).first()
        if not row:
            return {"success": False, "message": "数据源关系不存在"}

        relation, datasource = row
        relation.is_deleted = 1

        self.session.commit()
        return {"success": True, "message": "移除成功"}

    def _get_conversation(self, insight_conversation_id: int) -> InsightNsConversation | None:
        return self.session.query(InsightNsConversation).filter(
            InsightNsConversation.id == insight_conversation_id,
            InsightNsConversation.is_deleted == 0,
        ).first()

    def _normalize_uns_nodes(
        self,
        nodes: list[dict[str, Any]] | None,
    ) -> list[dict[str, Any]]:
        normalized: list[dict[str, Any]] = []
        seen: set[str] = set()

        for node in nodes or []:
            if not isinstance(node, dict):
                continue
            node_id = str(node.get('id') or node.get('uns_node_id') or '').strip()
            if not node_id or node_id in seen:
                continue
            count_children = int(node.get('countChildren') or node.get('count_children') or 0)
            is_folder = bool(node.get('isFolder') or node.get('is_folder'))
            if not is_folder:
                is_folder = bool(node.get('hasChildren')) or count_children > 0
            data_type = self._normalize_uns_data_type(node)
            normalized.append({
                "id": node_id,
                "alias": str(node.get('alias') or '').strip(),
                "name": str(node.get('name') or node.get('label') or node.get('pathName') or node_id).strip(),
                "data_type": data_type,
                "uns_node_path": self._normalize_uns_node_id_list(
                    node.get('unsNodePath')
                    or node.get('uns_node_path')
                    or node.get('ancestorNodeIds')
                    or node.get('ancestor_node_ids')
                    or []
                ),
                "count_children": count_children,
                "is_folder": is_folder,
            })
            seen.add(node_id)
        return normalized

    def _expand_uns_file_nodes(
        self,
        selected_nodes: list[dict[str, Any]],
        authorization: str,
    ) -> dict[str, list[dict[str, Any]]]:
        file_nodes: list[dict[str, Any]] = []
        selections: list[dict[str, Any]] = []
        failed: list[dict[str, Any]] = []
        seen_file_ids: set[str] = set()

        for node in selected_nodes:
            node_id = node["id"]
            if not node.get("is_folder"):
                data_type = node.get("data_type")
                if data_type is None:
                    failed.append({
                        "id": node_id,
                        "name": node.get("name", ""),
                        "message": "UNS 节点缺少 dataType，无法判断数据源类型",
                    })
                    continue
                if not self._is_supported_uns_table_node(node):
                    failed.append({
                        "id": node_id,
                        "name": node.get("name", ""),
                        "message": f"暂不支持导入 dataType={data_type} 的 UNS 节点",
                    })
                    continue
                if node_id not in seen_file_ids:
                    file_nodes.append(node)
                    seen_file_ids.add(node_id)
                selections.append({
                    **node,
                    "expanded_uns_node_ids": [node_id],
                })
                continue

            try:
                if int(node.get("count_children") or 0) > UNS_MAX_EXPANDED_FILES:
                    raise ValueError(
                        f"UNS 文件夹“{node.get('name') or node_id}”子节点过多（{node.get('count_children')}），"
                        f"请继续展开后选择更小范围（当前上限 {UNS_MAX_EXPANDED_FILES}）"
                    )
                expanded = self._expand_uns_folder_files(node, authorization)
                for file_node in expanded:
                    file_id = file_node["id"]
                    if file_id in seen_file_ids:
                        continue
                    file_nodes.append(file_node)
                    seen_file_ids.add(file_id)
                selections.append({
                    **node,
                    "expanded_uns_node_ids": [item["id"] for item in expanded],
                })
            except Exception as exc:
                failed.append({
                    "id": node_id,
                    "name": node.get("name", ""),
                    "message": str(exc),
                })

        return {
            "file_nodes": file_nodes,
            "selections": selections,
            "failed": failed,
        }

    def _expand_uns_folder_files(self, root_node: dict[str, Any], authorization: str) -> list[dict[str, Any]]:
        files: list[dict[str, Any]] = []
        stack: list[tuple[dict[str, Any], int]] = [(root_node, 0)]

        while stack:
            folder, depth = stack.pop()
            if depth >= UNS_MAX_EXPAND_DEPTH:
                raise ValueError(f"UNS 文件夹展开深度超过限制 {UNS_MAX_EXPAND_DEPTH}")

            for child in self._fetch_all_uns_children(folder["id"], authorization):
                child_id = str(child.get('id') or child.get('alias') or '').strip()
                if not child_id:
                    continue
                child_node = {
                    "id": child_id,
                    "alias": str(child.get('alias') or '').strip(),
                    "name": str(child.get('name') or child.get('pathName') or child_id).strip(),
                    "path": str(child.get('pathName') or child.get('path') or '').strip(),
                    "data_type": self._normalize_uns_data_type(child),
                    "is_folder": bool(child.get('hasChildren'))
                    or int(child.get('countChildren') or 0) > 0
                    or int(child.get('type') if child.get('type') is not None else -1) == 0,
                }
                if child_node["is_folder"]:
                    stack.append((child_node, depth + 1))
                    continue
                if not self._is_supported_uns_table_node(child_node):
                    continue

                files.append(child_node)
                if len(files) > UNS_MAX_EXPANDED_FILES:
                    raise ValueError(f"UNS 文件夹展开文件数超过限制 {UNS_MAX_EXPANDED_FILES}")

        return files

    def _fetch_all_uns_children(self, parent_id: str, authorization: str) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = []
        page_no = 1
        total = None

        while True:
            payload = supos_kernel_api.fetch_uns_tree_nodes(
                authorization=authorization,
                parent_id=parent_id,
                page_no=page_no,
                page_size=UNS_TREE_PAGE_SIZE,
                keyword='',
                search_type=1,
                data_type=sorted(UNS_TABLE_DATA_TYPES),
            )
            page_rows = payload.get("data") or []
            rows.extend(page_rows)
            total = int(payload.get("total") or len(rows)) if total is None else total

            if not page_rows or len(rows) >= total:
                break
            page_no += 1
            if page_no > 100:
                raise ValueError("UNS 文件夹分页过多，已停止展开")

        return rows

    def _normalize_uns_data_type(self, node: dict[str, Any]) -> int | None:
        raw_value = node.get('dataType')
        if raw_value is None:
            raw_value = node.get('data_type')
        if raw_value is None:
            return None
        try:
            return int(raw_value)
        except (TypeError, ValueError):
            return None

    def _is_supported_uns_table_node(self, node: dict[str, Any]) -> bool:
        data_type = node.get("data_type")
        return data_type in UNS_TABLE_DATA_TYPES

    def _fetch_uns_details_in_parallel(
        self,
        ids: list[str],
        authorization: str,
    ) -> list[dict[str, Any]]:
        if not ids:
            return []

        results: list[dict[str, Any]] = []
        with ThreadPoolExecutor(max_workers=min(UNS_DETAIL_WORKERS, max(1, len(ids)))) as executor:
            future_to_id = {
                executor.submit(supos_kernel_api.fetch_uns_instance_detail, node_id, authorization): node_id
                for node_id in ids
            }
            for future in as_completed(future_to_id):
                node_id = future_to_id[future]
                try:
                    detail = future.result()
                    if not str(detail.get('id') or '').strip():
                        detail = {**detail, "id": node_id}
                    results.append({
                        "success": True,
                        "id": node_id,
                        "detail": detail,
                    })
                except Exception as exc:
                    results.append({
                        "success": False,
                        "id": node_id,
                        "message": str(exc),
                    })
        results.sort(key=lambda item: ids.index(item["id"]))
        return results

    def _upsert_uns_table_datasource(
        self,
        detail: dict[str, Any],
        lake_rds_database_name: str,
        uns_data_type: int | None = None,
        ts_query_source: str = '',
        reference_detail_by_id: dict[str, dict[str, Any]] | None = None,
    ) -> InsightDatasource:
        uns_node_id = str(detail.get('id') or '').strip()
        if not uns_node_id:
            raise ValueError('UNS 节点缺少 id，无法导入')
        alias = str(detail.get('alias') or '').strip()
        if not alias:
            raise ValueError('UNS 节点缺少 alias，无法导入')

        datasource = self._find_uns_datasource_by_node_id(uns_node_id)
        datasource_name = str(detail.get('name') or alias).strip()[:128] or alias

        if datasource is None:
            datasource_name = self._build_unique_datasource_name(SHARED_UNS_NAMESPACE_ID, datasource_name)
            datasource = InsightDatasource(
                insight_namespace_id=SHARED_UNS_NAMESPACE_ID,
                datasource_type='table',
                datasource_name=datasource_name,
                knowledge_tag='',
                uns_node_id=uns_node_id,
                datasource_schema='',
                datasource_config_json='{}',
            )
            self.session.add(datasource)
            self.session.flush()
        else:
            datasource.datasource_name = datasource_name
            datasource.uns_node_id = uns_node_id

        data_type = (
            self._normalize_uns_data_type({'data_type': uns_data_type})
            if uns_data_type is not None
            else self._normalize_uns_data_type(detail)
        )
        if data_type in {UNS_TIME_SERIES_DATA_TYPE, UNS_CALCULATED_TABLE_DATA_TYPE}:
            datasource.datasource_schema = self._build_uns_time_series_datasource_schema(
                detail,
                datasource.datasource_name,
                ts_query_source=ts_query_source,
            )
            datasource.datasource_config_json = dump_json(
                self._build_time_series_table_config(
                    detail=detail,
                    alias=alias,
                    lake_rds_database_name=lake_rds_database_name,
                    ts_query_source=ts_query_source,
                    uns_data_type=data_type,
                )
            )
        elif data_type == UNS_PROTOCOL_TABLE_DATA_TYPE:
            datasource.datasource_schema = self._build_uns_datasource_schema(detail, datasource.datasource_name)
            datasource.datasource_config_json = dump_json(
                self._build_protocol_table_config(detail=detail, alias=alias)
            )
        elif data_type == UNS_REFERENCE_TABLE_DATA_TYPE:
            ref_info = self._extract_reference_uns_ref(detail)
            ref_id = str(ref_info.get('id') or '').strip()
            ref_alias = str(ref_info.get('alias') or '').strip()
            ref_detail = (reference_detail_by_id or {}).get(ref_id) if ref_id else None
            if not isinstance(ref_detail, dict) or not ref_detail:
                ref_detail = dict(ref_info)
            ref_data_type = self._normalize_uns_data_type(ref_detail) or self._normalize_uns_data_type(ref_info)

            if ref_data_type == UNS_TIME_SERIES_DATA_TYPE:
                if not normalize_ts_query_source(ts_query_source):
                    raise ValueError('引用节点类型暂不支持')
                datasource.datasource_schema = self._build_uns_time_series_datasource_schema(
                    ref_detail,
                    datasource.datasource_name,
                    ts_query_source=ts_query_source,
                )
                config = self._build_time_series_table_config(
                    detail=ref_detail,
                    alias=ref_alias,
                    lake_rds_database_name=lake_rds_database_name,
                    ts_query_source=ts_query_source,
                )
                config.update(self._build_reference_config_overlay(detail, alias, ref_info, ref_data_type))
                datasource.datasource_config_json = dump_json(config)
            elif ref_data_type == UNS_RELATIONAL_TABLE_DATA_TYPE:
                datasource.datasource_schema = self._build_uns_datasource_schema(ref_detail, datasource.datasource_name)
                datasource.datasource_config_json = dump_json(
                    self._build_reference_relational_table_config(
                        detail=detail,
                        alias=alias,
                        ref_info=ref_info,
                        lake_rds_database_name=lake_rds_database_name,
                    )
                )
            else:
                raise ValueError('引用节点类型暂不支持')
        else:
            datasource.datasource_schema = self._build_uns_datasource_schema(detail, datasource.datasource_name)
            datasource.datasource_config_json = dump_json(
                self._build_relational_table_config(
                    detail=detail,
                    alias=alias,
                    data_type=data_type,
                    lake_rds_database_name=lake_rds_database_name,
                )
            )
        return datasource

    def _fetch_reference_uns_details_for_import(
        self,
        details: list[dict[str, Any]],
        authorization: str,
    ) -> tuple[dict[str, dict[str, Any]], list[dict[str, Any]]]:
        reference_ids: list[str] = []
        available_detail_by_id: dict[str, dict[str, Any]] = {
            str(detail.get('id') or '').strip(): detail
            for detail in details
            if str(detail.get('id') or '').strip()
        }
        reference_detail_by_id: dict[str, dict[str, Any]] = {}
        owner_by_ref_id: dict[str, str] = {}
        failed_rows: list[dict[str, Any]] = []

        for detail in details:
            if self._normalize_uns_data_type(detail) != UNS_REFERENCE_TABLE_DATA_TYPE:
                continue
            owner_id = str(detail.get('id') or '').strip()
            try:
                ref_info = self._extract_reference_uns_ref(detail)
            except ValueError as exc:
                failed_rows.append({"id": owner_id, "message": str(exc)})
                continue

            ref_id = str(ref_info.get('id') or '').strip()
            if not ref_id:
                if self._normalize_uns_data_type(ref_info) not in {
                    UNS_RELATIONAL_TABLE_DATA_TYPE,
                    UNS_TIME_SERIES_DATA_TYPE,
                }:
                    failed_rows.append({"id": owner_id, "message": "引用节点类型暂不支持"})
                continue
            if ref_id in available_detail_by_id:
                ref_data_type = self._normalize_uns_data_type(available_detail_by_id[ref_id])
                if ref_data_type not in {UNS_RELATIONAL_TABLE_DATA_TYPE, UNS_TIME_SERIES_DATA_TYPE}:
                    failed_rows.append({"id": owner_id, "message": "引用节点类型暂不支持"})
                else:
                    reference_detail_by_id[ref_id] = available_detail_by_id[ref_id]
                continue
            if ref_id not in owner_by_ref_id:
                owner_by_ref_id[ref_id] = owner_id
                reference_ids.append(ref_id)

        if not reference_ids or failed_rows:
            return {}, failed_rows

        for result in self._fetch_uns_details_in_parallel(reference_ids, authorization):
            ref_id = str(result.get("id") or (result.get("detail") or {}).get("id") or "").strip()
            owner_id = owner_by_ref_id.get(ref_id, ref_id)
            if not result.get("success"):
                failed_rows.append({
                    "id": owner_id,
                    "message": result.get("message", "引用节点详情同步失败"),
                })
                continue

            ref_detail = result.get("detail") if isinstance(result.get("detail"), dict) else {}
            ref_data_type = self._normalize_uns_data_type(ref_detail)
            if ref_data_type not in {UNS_RELATIONAL_TABLE_DATA_TYPE, UNS_TIME_SERIES_DATA_TYPE}:
                failed_rows.append({"id": owner_id, "message": "引用节点类型暂不支持"})
                continue
            reference_detail_by_id[ref_id] = ref_detail

        return reference_detail_by_id, failed_rows

    def _build_relational_table_config(
        self,
        *,
        detail: dict[str, Any],
        alias: str,
        data_type: int | None,
        lake_rds_database_name: str,
    ) -> dict[str, Any]:
        table_alias = alias
        payload: dict[str, Any] = {}
        if data_type == UNS_REFERENCE_TABLE_DATA_TYPE:
            table_alias = self._extract_reference_uns_alias(detail)
            payload["refer_alias"] = table_alias

        payload.update({
            "database_name": lake_rds_database_name,
            "table_name": f"public.{table_alias}",
            "uns_alias": alias,
            "uns_data_type": data_type or UNS_RELATIONAL_TABLE_DATA_TYPE,
            "table_engine": TABLE_ENGINE_RELATIONAL,
            "uns_path": detail.get('path') or '',
            "uns_path_name": detail.get('pathName') or detail.get('name') or '',
        })
        return payload

    def _build_reference_relational_table_config(
        self,
        *,
        detail: dict[str, Any],
        alias: str,
        ref_info: dict[str, Any],
        lake_rds_database_name: str,
    ) -> dict[str, Any]:
        ref_alias = str(ref_info.get('alias') or '').strip()
        payload = {
            "database_name": lake_rds_database_name,
            "table_name": f"public.{ref_alias}",
            "table_engine": TABLE_ENGINE_RELATIONAL,
        }
        payload.update(self._build_reference_config_overlay(
            detail,
            alias,
            ref_info,
            UNS_RELATIONAL_TABLE_DATA_TYPE,
        ))
        return payload

    def _build_reference_config_overlay(
        self,
        detail: dict[str, Any],
        alias: str,
        ref_info: dict[str, Any],
        ref_data_type: int | None,
    ) -> dict[str, Any]:
        return {
            "uns_data_type": UNS_REFERENCE_TABLE_DATA_TYPE,
            "source_uns_alias": alias,
            "refer_node_id": str(ref_info.get('id') or '').strip(),
            "refer_alias": str(ref_info.get('alias') or '').strip(),
            "refer_data_type": ref_data_type,
            "uns_path": detail.get('path') or '',
            "uns_path_name": detail.get('pathName') or detail.get('name') or '',
        }

    def _extract_reference_uns_alias(self, detail: dict[str, Any]) -> str:
        return str(self._extract_reference_uns_ref(detail).get('alias') or '').strip()

    def _extract_reference_uns_ref(self, detail: dict[str, Any]) -> dict[str, Any]:
        refers = detail.get('refers') or []
        if not refers:
            refers = detail.get('refer') or []
        if isinstance(refers, str):
            refers = safe_json_loads(refers, [])
        if isinstance(refers, dict):
            refers = [refers]
        if not isinstance(refers, list) or not refers:
            raise ValueError('dataType=7 的 UNS 节点 refers 缺失或为空')
        first_ref = refers[0]
        if not isinstance(first_ref, dict):
            raise ValueError('dataType=7 的 UNS 节点 refers[0] 无效')
        ref_alias = str(first_ref.get('alias') or '').strip()
        if not ref_alias:
            raise ValueError('dataType=7 的 UNS 节点 refers[0].alias 缺失或无效')
        return dict(first_ref)

    def _build_protocol_table_config(self, *, detail: dict[str, Any], alias: str) -> dict[str, Any]:
        protocol = safe_json_loads(detail.get('protocol'), {})
        db_protocol = protocol.get('db') if isinstance(protocol, dict) else {}
        if not isinstance(db_protocol, dict):
            db_protocol = {}

        protocol_alias = str(db_protocol.get('alias') or '').strip()
        schema_name = str(db_protocol.get('schemaName') or '').strip()
        table_name = str(db_protocol.get('tableName') or '').strip()
        if not protocol_alias:
            raise ValueError('dataType=9 的 UNS 节点 protocol.db.alias 缺失或无效')
        if not schema_name or not table_name:
            raise ValueError('dataType=9 的 UNS 节点 protocol.db.schemaName/tableName 缺失')

        return {
            "database_name": f"fqe_{protocol_alias}",
            "table_name": f"{schema_name}.{table_name}",
            "uns_alias": alias,
            "uns_data_type": UNS_PROTOCOL_TABLE_DATA_TYPE,
            "table_engine": TABLE_ENGINE_RELATIONAL,
            "uns_path": detail.get('path') or '',
            "uns_path_name": detail.get('pathName') or detail.get('name') or '',
        }

    def _build_time_series_table_config(
        self,
        *,
        detail: dict[str, Any],
        alias: str,
        lake_rds_database_name: str,
        ts_query_source: str,
        uns_data_type: int = UNS_TIME_SERIES_DATA_TYPE,
    ) -> dict[str, Any]:
        normalized_ts_query_source = normalize_ts_query_source(ts_query_source)
        time_field = 'timeStamp'
        value_fields = self._extract_time_series_value_fields(
            detail,
            excluded_fields={'timeStamp', 'status'},
        )
        payload = {
            "database_name": lake_rds_database_name,
            "table_name": alias,
            "uns_alias": alias,
            "uns_data_type": uns_data_type,
            "table_engine": TABLE_ENGINE_TIME_SERIES,
            "ts_query_source": normalized_ts_query_source,
            "time_field": time_field,
            "value_fields": value_fields or ["value"],
            "quality_field": "status",
            "uns_path": detail.get('path') or '',
            "uns_path_name": detail.get('pathName') or detail.get('name') or '',
        }
        return payload

    def _extract_time_series_value_fields(
        self,
        detail: dict[str, Any],
        excluded_fields: set[str],
    ) -> list[str]:
        excluded = {str(item or '').strip().lower() for item in excluded_fields if str(item or '').strip()}
        value_fields: list[str] = []
        seen: set[str] = set()
        for field in detail.get('fields') or []:
            field_name = str(field.get('name') or '').strip()
            if not field_name:
                continue
            normalized = field_name.lower()
            if normalized in excluded or normalized in seen:
                continue
            value_fields.append(field_name)
            seen.add(normalized)
        return value_fields

    def _find_uns_datasource_by_node_id(
        self,
        uns_node_id: str,
    ) -> InsightDatasource | None:
        return self.session.query(InsightDatasource).filter(
            InsightDatasource.uns_node_id == uns_node_id,
            InsightDatasource.datasource_type == 'table',
            InsightDatasource.is_deleted == 0,
        ).first()

    def _find_datasource_ids_by_uns_node_ids(self, uns_node_ids: list[str]) -> list[int]:
        normalized_ids = [
            str(item or '').strip()
            for item in uns_node_ids or []
            if str(item or '').strip()
        ]
        if not normalized_ids:
            return []

        rows = self.session.query(InsightDatasource.id).filter(
            InsightDatasource.uns_node_id.in_(normalized_ids),
            InsightDatasource.datasource_type == 'table',
            InsightDatasource.is_deleted == 0,
        ).all()
        return [int(row[0]) for row in rows]

    def _soft_delete_conversation_datasource_relations(
        self,
        insight_conversation_id: int,
        datasource_ids: list[int],
    ) -> None:
        if not datasource_ids:
            return
        self.session.query(InsightNsRelDatasource).filter(
            InsightNsRelDatasource.insight_conversation_id == insight_conversation_id,
            InsightNsRelDatasource.datasource_id.in_(datasource_ids),
            InsightNsRelDatasource.is_deleted == 0,
        ).update(
            {InsightNsRelDatasource.is_deleted: 1},
            synchronize_session=False,
        )

    def _namespace_has_same_named_datasource(self, insight_namespace_id: int, datasource_name: str) -> bool:
        existing = self.session.query(InsightDatasource.id).filter(
            InsightDatasource.insight_namespace_id == insight_namespace_id,
            InsightDatasource.datasource_name == datasource_name,
            InsightDatasource.is_deleted == 0,
        ).first()
        return existing is not None

    def _build_unique_datasource_name(self, insight_namespace_id: int, base_name: str) -> str:
        normalized_base = (base_name or '外部数据源').strip()[:128] or '外部数据源'
        if not self._namespace_has_same_named_datasource(insight_namespace_id, normalized_base):
            return normalized_base

        suffix = 2
        while True:
            candidate = f"{normalized_base}_{suffix}"[:128]
            if not self._namespace_has_same_named_datasource(insight_namespace_id, candidate):
                return candidate
            suffix += 1

    def _bind_datasource(
        self,
        insight_namespace_id: int,
        insight_conversation_id: int,
        datasource_id: int,
        bind_source: str,
    ) -> InsightNsRelDatasource:
        relation = self.session.query(InsightNsRelDatasource).filter(
            InsightNsRelDatasource.insight_conversation_id == insight_conversation_id,
            InsightNsRelDatasource.datasource_id == datasource_id,
        ).first()
        if relation is not None:
            relation.insight_namespace_id = insight_namespace_id
            relation.is_deleted = 0
            relation.is_active = 1
            relation.bind_source = bind_source
            if not relation.sort_no:
                relation.sort_no = self._next_sort_no(insight_conversation_id)
            return relation

        relation = InsightNsRelDatasource(
            insight_namespace_id=insight_namespace_id,
            insight_conversation_id=insight_conversation_id,
            datasource_id=datasource_id,
            is_active=1,
            sort_no=self._next_sort_no(insight_conversation_id),
            bind_source=bind_source,
        )
        self.session.add(relation)
        self.session.flush()
        return relation

    def _upsert_uns_selection(
        self,
        insight_namespace_id: int,
        insight_conversation_id: int,
        selection: dict[str, Any],
    ) -> InsightNsUnsSelection:
        row = self.session.query(InsightNsUnsSelection).filter(
            InsightNsUnsSelection.insight_conversation_id == insight_conversation_id,
            InsightNsUnsSelection.uns_node_id == selection["id"],
        ).first()
        if row is None:
            row = InsightNsUnsSelection(
                insight_namespace_id=insight_namespace_id,
                insight_conversation_id=insight_conversation_id,
                uns_node_id=selection["id"],
            )
            self.session.add(row)

        row.insight_namespace_id = insight_namespace_id
        row.uns_node_name = selection.get("name", "")[:255]
        row.uns_node_path = self._build_uns_selection_path_payload(selection.get("uns_node_path") or [])
        row.is_folder = 1 if selection.get("is_folder") else 0
        row.expanded_uns_node_ids_json = dump_json(selection.get("expanded_uns_node_ids") or [])
        row.is_deleted = 0
        self.session.flush()
        return row

    def _canonicalize_uns_selections(
        self,
        selections: list[dict[str, Any]],
    ) -> list[dict[str, Any]]:
        normalized_selections: list[dict[str, Any]] = []
        for selection in selections or []:
            selection_id = str(selection.get("id") or "").strip()
            if not selection_id:
                continue
            normalized_selections.append({
                **selection,
                "id": selection_id,
                "expanded_uns_node_ids": self._normalize_uns_node_id_list(selection.get("expanded_uns_node_ids") or [selection_id]),
            })

        sorted_selections = sorted(
            normalized_selections,
            key=lambda item: (
                0 if item.get("is_folder") else 1,
                -len(item.get("expanded_uns_node_ids") or []),
                item["id"],
            ),
        )
        kept: list[dict[str, Any]] = []
        covered_node_ids: set[str] = set()
        for selection in sorted_selections:
            selection_id = selection["id"]
            expanded_ids = selection.get("expanded_uns_node_ids") or [selection_id]
            if selection_id in covered_node_ids or set(expanded_ids).issubset(covered_node_ids):
                continue
            kept.append(selection)
            covered_node_ids.update(expanded_ids)
            covered_node_ids.add(selection_id)
        return kept

    def _normalize_uns_node_id_list(self, values: list[Any]) -> list[str]:
        normalized_ids: list[str] = []
        seen: set[str] = set()
        for item in values:
            node_id = str(item or '').strip()
            if not node_id or node_id in seen:
                continue
            normalized_ids.append(node_id)
            seen.add(node_id)
        return normalized_ids

    def _build_uns_selection_path_payload(
        self,
        uns_node_path: list[Any],
    ) -> str:
        return dump_json(self._normalize_uns_node_id_list(uns_node_path or []))

    def _upsert_uns_file_selection_by_node_id(
        self,
        insight_namespace_id: int,
        insight_conversation_id: int,
        uns_node_id: str,
    ) -> InsightNsUnsSelection:
        datasource = self._find_uns_datasource_by_node_id(uns_node_id)
        datasource_config = safe_json_loads(datasource.datasource_config_json, {}) if datasource else {}
        return self._upsert_uns_selection(
            insight_namespace_id=insight_namespace_id,
            insight_conversation_id=insight_conversation_id,
            selection={
                "id": uns_node_id,
                "name": datasource.datasource_name if datasource else uns_node_id,
                "uns_node_path": [],
                "is_folder": False,
                "expanded_uns_node_ids": [uns_node_id],
            },
        )

    def _replace_uns_selections(
        self,
        insight_namespace_id: int,
        insight_conversation_id: int,
        selections: list[dict[str, Any]],
    ) -> None:
        self.session.query(InsightNsUnsSelection).filter(
            InsightNsUnsSelection.insight_conversation_id == insight_conversation_id,
            InsightNsUnsSelection.is_deleted == 0,
        ).update(
            {InsightNsUnsSelection.is_deleted: 1},
            synchronize_session=False,
        )
        self.session.flush()

        for selection in selections:
            self._upsert_uns_selection(
                insight_namespace_id=insight_namespace_id,
                insight_conversation_id=insight_conversation_id,
                selection=selection,
            )

    def _find_uns_datasources_by_node_ids(
        self,
        uns_node_ids: list[str],
    ) -> dict[str, InsightDatasource]:
        normalized_ids = self._normalize_uns_node_id_list(uns_node_ids)
        if not normalized_ids:
            return {}

        rows = self.session.query(InsightDatasource).filter(
            InsightDatasource.uns_node_id.in_(normalized_ids),
            InsightDatasource.datasource_type == 'table',
            InsightDatasource.is_deleted == 0,
        ).all()
        return {
            str(row.uns_node_id or '').strip(): row
            for row in rows
            if str(row.uns_node_id or '').strip()
        }

    def _sync_uns_conversation_bindings(
        self,
        insight_namespace_id: int,
        insight_conversation_id: int,
        desired_datasource_ids: list[int],
    ) -> list[int]:
        desired_ids = sorted(set(int(item) for item in desired_datasource_ids or [] if int(item) > 0))
        active_rows = self.session.query(InsightNsRelDatasource.id, InsightNsRelDatasource.datasource_id).join(
            InsightDatasource,
            InsightDatasource.id == InsightNsRelDatasource.datasource_id,
        ).filter(
            InsightNsRelDatasource.insight_conversation_id == insight_conversation_id,
            InsightNsRelDatasource.bind_source == BIND_SOURCE_IMPORT_SELECTED,
            InsightNsRelDatasource.is_deleted == 0,
            InsightDatasource.is_deleted == 0,
            InsightDatasource.uns_node_id != '',
        ).all()

        active_ids = {int(row[1]) for row in active_rows}
        removable_ids = [datasource_id for datasource_id in active_ids if datasource_id not in desired_ids]
        self._soft_delete_conversation_datasource_relations(
            insight_conversation_id=insight_conversation_id,
            datasource_ids=removable_ids,
        )

        for datasource_id in desired_ids:
            self._bind_datasource(
                insight_namespace_id=insight_namespace_id,
                insight_conversation_id=insight_conversation_id,
                datasource_id=datasource_id,
                bind_source=BIND_SOURCE_IMPORT_SELECTED,
            )
        return removable_ids

    def _next_sort_no(self, insight_conversation_id: int) -> int:
        count = self.session.query(InsightNsRelDatasource.id).filter(
            InsightNsRelDatasource.insight_conversation_id == insight_conversation_id,
            InsightNsRelDatasource.is_deleted == 0,
        ).count()
        return count + 1


    def _build_uns_datasource_schema(self, detail: dict[str, Any], datasource_name: str) -> str:
        definition = detail.get('fields') or []
        properties: dict[str, PropertySchema] = {}
        required: list[str] = []

        for field in definition:
            field_name = str(field.get('name') or '').strip()
            if not field_name:
                continue

            description_parts = [
                str(field.get('displayName') or '').strip(),
                str(field.get('remark') or '').strip(),
            ]
            properties[field_name] = PropertySchema(
                type=self._map_uns_field_type(field.get('type')),
                description='；'.join(part for part in description_parts if part) or field_name,
                example=None,
            )
            if not bool(field.get('systemField')):
                required.append(field_name)

        schema = DataSourceSchema(
            name=datasource_name,
            description=str(detail.get('description') or '').strip()
            or f"UNS 节点“{detail.get('name') or datasource_name}”转换得到的表结构",
            properties=properties,
            required=required,
        )
        return dump_json(schema.model_dump())

    def _build_file_preview_summary(self, dataframe: pd.DataFrame, sheet_name: str | None = None) -> dict[str, Any]:
        sample_frame = dataframe.head(20)
        columns = [str(column) for column in dataframe.columns]
        sample_rows = [
            {
                str(column): self._preview_value(row.get(column))
                for column in dataframe.columns
            }
            for row in sample_frame.to_dict(orient='records')
        ]

        column_stats: dict[str, Any] = {}
        for column in dataframe.columns:
            series = dataframe[column]
            stats: dict[str, Any] = {
                "non_null_count": int(series.notna().sum()),
                "null_count": int(series.isna().sum()),
                "example": self._preview_value(self._extract_example_value(series)),
            }
            if pd.api.types.is_numeric_dtype(series):
                numeric = pd.to_numeric(series, errors='coerce').dropna()
                if not numeric.empty:
                    stats.update({
                        "min": self._preview_value(numeric.min()),
                        "max": self._preview_value(numeric.max()),
                        "mean": self._preview_value(numeric.mean()),
                    })
            else:
                values = series.dropna().astype(str)
                if not values.empty:
                    stats["unique_preview_count"] = int(values.nunique(dropna=True))
                    stats["top_values"] = [
                        self._preview_value(value)
                        for value in values.value_counts().head(5).index.tolist()
                    ]
            column_stats[str(column)] = stats

        return {
            "sheet_name": sheet_name or "",
            "columns": columns,
            "row_count_preview": int(len(dataframe)),
            "sample_rows": sample_rows,
            "column_stats": column_stats,
        }

    def _preview_value(self, value: Any) -> Any:
        if value is None:
            return None
        try:
            if pd.isna(value):
                return None
        except Exception:
            pass
        if hasattr(value, 'isoformat'):
            try:
                return value.isoformat()
            except Exception:
                return str(value)
        if isinstance(value, (int, float, bool)):
            return value
        text = str(value)
        return text if len(text) <= 120 else f"{text[:117]}..."

    def _build_uns_time_series_datasource_schema(
        self,
        detail: dict[str, Any],
        datasource_name: str,
        ts_query_source: str,
    ) -> str:
        normalized_ts_query_source = normalize_ts_query_source(ts_query_source)
        base_payload = safe_json_loads(self._build_uns_datasource_schema(detail, datasource_name), {})
        properties = dict(base_payload.get('properties') or {})
        required = list(base_payload.get('required') or [])

        def add_property(name: str, type_name: str, description: str) -> None:
            if not name or name in properties:
                return
            properties[name] = PropertySchema(
                type=type_name,
                description=description,
                example=None,
            ).model_dump()
            if name not in required:
                required.append(name)

        if normalized_ts_query_source == 'timescale':
            description = (
                f"时序数据源“{detail.get('name') or datasource_name}”，ts_query_source 为 timescale。"
                "查询时应通过 ts_query('timescale', '<内部 SQL>') 访问。"
            )
        else:
            description = (
                f"时序数据源“{detail.get('name') or datasource_name}”，ts_query_source 为 vxbase。"
                "查询时应通过 ts_query('vxbase', '<内部 SQL>') 访问。"
            )
        add_property('timeStamp', 'string', '时序数据时间字段')
        add_property('status', 'integer', '时序数据质量码，通常 0 表示好值')

        schema = DataSourceSchema(
            name=datasource_name,
            description=description,
            properties={
                key: value if isinstance(value, PropertySchema) else PropertySchema(**value)
                for key, value in properties.items()
                if isinstance(value, (dict, PropertySchema))
            },
            required=required,
        )
        return dump_json(schema.model_dump())


    def _read_csv_preview(self, file_path: Path) -> pd.DataFrame:
        """优先自动尝试常见 CSV 编码，减少编码差异导致的上传失败。"""
        candidate_encodings = (
            'utf-8',
            'utf-8-sig',
            'gb18030',
            'gbk',
        )
        last_decode_error: UnicodeDecodeError | None = None

        for encoding in candidate_encodings:
            try:
                return pd.read_csv(file_path, nrows=50, encoding=encoding)
            except UnicodeDecodeError as exc:
                last_decode_error = exc
                continue

        if last_decode_error is not None:
            raise ValueError(
                'CSV 文件编码解析失败。系统已自动尝试 UTF-8、UTF-8-SIG、GB18030、GBK，'
                '仍无法读取，请将文件另存为 UTF-8 编码后重试。'
            ) from last_decode_error

        # 不是编码问题时，保留原始异常语义，避免误导用户。
        return pd.read_csv(file_path, nrows=50)

    def _infer_schema_type(self, series: pd.Series) -> str:
        dtype = str(series.dtype).lower()
        if 'int' in dtype:
            return 'integer'
        if 'float' in dtype or 'double' in dtype or 'decimal' in dtype:
            return 'number'
        if 'bool' in dtype:
            return 'boolean'
        if 'date' in dtype or 'time' in dtype:
            return 'string'
        return 'string'

    def _map_uns_field_type(self, source_type: Any) -> str:
        normalized = str(source_type or '').strip().upper()
        if normalized in {'INTEGER', 'LONG', 'INT', 'SHORT'}:
            return 'integer'
        if normalized in {'DOUBLE', 'FLOAT', 'DECIMAL', 'NUMBER'}:
            return 'number'
        if normalized in {'BOOLEAN', 'BOOL'}:
            return 'boolean'
        return 'string'

    def _extract_example_value(self, series: pd.Series) -> Any | None:
        non_null = series.dropna()
        if non_null.empty:
            return None
        value = non_null.iloc[0]
        if hasattr(value, 'isoformat'):
            try:
                return value.isoformat()
            except Exception:
                return str(value)
        if isinstance(value, (int, float, bool)):
            return value
        return str(value)

    def _get_or_create_datasource(
        self,
        insight_namespace_id: int,
        datasource_type: str,
        datasource_name: str,
        knowledge_tag: str,
        datasource_schema: str,
        datasource_config_json: str,
    ) -> InsightDatasource:
        """按空间、名称和类型幂等创建数据源定义。"""
        normalized_type = normalize_datasource_type(datasource_type)
        if normalized_type == 'unknown':
            raise ValueError(f'不支持的数据源类型: {datasource_type}')

        datasource = self.session.query(InsightDatasource).filter(
            InsightDatasource.insight_namespace_id == insight_namespace_id,
            InsightDatasource.datasource_name == datasource_name,
            InsightDatasource.datasource_type == normalized_type,
            InsightDatasource.is_deleted == 0,
        ).first()
        if datasource is not None:
            datasource.knowledge_tag = knowledge_tag or datasource.knowledge_tag or ''
            datasource.datasource_schema = datasource_schema or datasource.datasource_schema or ''
            datasource.datasource_config_json = datasource_config_json or datasource.datasource_config_json or '{}'
            return datasource

        datasource = InsightDatasource(
            insight_namespace_id=insight_namespace_id,
            datasource_type=normalized_type,
            datasource_name=datasource_name,
            knowledge_tag=knowledge_tag or '',
            datasource_schema=datasource_schema or '',
            datasource_config_json=datasource_config_json or '{}',
        )
        self.session.add(datasource)
        self.session.flush()
        return datasource


    def upload_file_datasource_to_namespace(
        self,
        insight_namespace_id: int,
        upload_file: FileStorage,
        upload_dir: str,
        selected_llm_model_id: str = '',
    ) -> dict[str, Any]:
        """上传文件并在空间下创建数据源；Excel 会按 sheet 拆成多个逻辑数据源。"""
        if insight_namespace_id <= 0:
            return {"success": False, "message": "空间不存在"}
        if upload_file is None or not upload_file.filename:
            return {"success": False, "message": "请选择要上传的文件"}

        original_filename = Path(upload_file.filename).name
        suffix = Path(original_filename).suffix.lower()
        if suffix not in {'.csv', '.xls', '.xlsx'}:
            return {"success": False, "message": "仅支持上传 csv、xls、xlsx 文件"}

        upload_dir_path = Path(upload_dir)
        upload_dir_path.mkdir(parents=True, exist_ok=True)
        stored_filename = f"{uuid4().hex}{suffix}"
        stored_path = upload_dir_path / stored_filename
        upload_file.save(str(stored_path))

        base_name = Path(original_filename).stem or f"外部数据源_{uuid4().hex[:6]}"
        try:
            datasource_payloads = self._build_uploaded_file_datasource_payloads(
                insight_namespace_id=insight_namespace_id,
                stored_path=stored_path,
                original_filename=original_filename,
                stored_filename=stored_filename,
                base_name=base_name,
            )
        except ValueError as exc:
            if stored_path.exists():
                stored_path.unlink(missing_ok=True)
            return {"success": False, "message": str(exc)}

        created_datasource_ids: list[int] = []
        description_items: list[dict[str, Any]] = []
        for payload in datasource_payloads:
            datasource = self._get_or_create_datasource(
                insight_namespace_id=insight_namespace_id,
                datasource_type='local_file',
                datasource_name=payload["datasource_name"],
                knowledge_tag=payload["knowledge_tag"],
                datasource_schema=payload["datasource_schema"],
                datasource_config_json=payload["datasource_config_json"],
            )
            created_datasource_ids.append(int(datasource.id or 0))
            description_items.append({
                "datasource_id": int(datasource.id or 0),
                "datasource_name": datasource.datasource_name,
                "original_filename": original_filename,
                "sheet_name": payload.get("sheet_name", ""),
                "preview_summary": payload.get("preview_summary") or {},
                "base_description": payload.get("base_description", ""),
                "schema_hash": payload.get("schema_hash", ""),
            })

        self.session.commit()
        self._submit_uploaded_datasource_description_job(
            description_items,
            selected_llm_model_id=selected_llm_model_id,
        )
        created_datasources = self._load_datasource_dicts(created_datasource_ids)
        message = "文件上传成功"
        if len(created_datasources) > 1:
            message = f"Excel 文件上传成功，已创建 {len(created_datasources)} 个工作表数据源"
        return {
            "success": True,
            "message": message,
            "data": self._build_upload_response_data(created_datasources),
        }

    def _submit_uploaded_datasource_description_job(
        self,
        description_items: list[dict[str, Any]],
        selected_llm_model_id: str = '',
    ) -> None:
        valid_items = [item for item in description_items if int(item.get("datasource_id") or 0) > 0]
        if not valid_items:
            return
        try:
            future = get_datasource_description_runner().submit(
                valid_items,
                selected_model_id=selected_llm_model_id,
            )
            wait_seconds = max(float(getattr(Config, 'DATASOURCE_DESCRIPTION_SYNC_WAIT_SECONDS', 2) or 0), 0)
            if wait_seconds > 0:
                future.result(timeout=wait_seconds)
                self.session.expire_all()
        except TimeoutError:
            logger.info("上传数据源描述仍在后台生成: datasource_ids=%s", [
                item.get("datasource_id") for item in valid_items
            ])
        except Exception:
            logger.warning("提交或等待上传数据源描述生成任务失败，保留默认描述", exc_info=True)

    def _load_datasource_dicts(self, datasource_ids: list[int]) -> list[dict[str, Any]]:
        normalized_ids = [int(item or 0) for item in datasource_ids if int(item or 0) > 0]
        if not normalized_ids:
            return []
        rows = self.session.query(InsightDatasource).filter(
            InsightDatasource.id.in_(normalized_ids),
            InsightDatasource.is_deleted == 0,
        ).all()
        row_map = {int(row.id): row for row in rows}
        return [
            self._datasource_to_dict(row_map[datasource_id])
            for datasource_id in normalized_ids
            if datasource_id in row_map
        ]

    def _build_file_datasource_schema(
        self,
        file_path: Path,
        datasource_name: str,
        sheet_name: str | None = None,
    ) -> str:
        """读取上传文件前几行数据，并推断元数据 Schema。"""
        datasource_schema, _preview_summary, _description = self._build_file_datasource_schema_with_preview(
            file_path,
            datasource_name,
            sheet_name=sheet_name,
        )
        return datasource_schema

    def _build_file_datasource_schema_with_preview(
        self,
        file_path: Path,
        datasource_name: str,
        sheet_name: str | None = None,
    ) -> tuple[str, dict[str, Any], str]:
        dataframe = self._read_file_preview(file_path, sheet_name=sheet_name)

        properties: dict[str, PropertySchema] = {}
        for column in dataframe.columns:
            series = dataframe[column]
            properties[str(column)] = PropertySchema(
                type=self._infer_schema_type(series),
                description=f"文件字段“{column}”",
                example=self._extract_example_value(series),
            )

        description = UPLOADED_DATASOURCE_DESCRIPTION_PENDING

        schema = DataSourceSchema(
            name=datasource_name,
            description=description,
            properties=properties,
            required=[str(column) for column in dataframe.columns],
        )
        return dump_json(schema.model_dump()), self._build_file_preview_summary(dataframe, sheet_name), description

    def _build_uploaded_file_datasource_payloads(
        self,
        insight_namespace_id: int,
        stored_path: Path,
        original_filename: str,
        stored_filename: str,
        base_name: str,
    ) -> list[dict[str, Any]]:
        suffix = stored_path.suffix.lower()
        file_path = str(stored_path).replace('\\', '/')
        file_size_bytes = stored_path.stat().st_size
        if suffix in {'.xlsx', '.xls'}:
            workbook_group_id = f"workbook_{uuid4().hex[:16]}"
            sheet_names = self._list_excel_sheet_names(stored_path)
            if not sheet_names:
                raise ValueError('Excel 文件未发现可用的工作表')

            payloads: list[dict[str, Any]] = []
            sheet_count = len(sheet_names)
            for index, raw_sheet_name in enumerate(sheet_names, start=1):
                sheet_name = str(raw_sheet_name or '').strip() or f"Sheet{index}"
                datasource_name = self._build_unique_datasource_name(
                    insight_namespace_id=insight_namespace_id,
                    base_name=f"{base_name} / {sheet_name}",
                )
                datasource_schema, preview_summary, base_description = self._build_file_datasource_schema_with_preview(
                    stored_path,
                    datasource_name,
                    sheet_name=sheet_name,
                )
                datasource_config_json = dump_json({
                    "file_path": file_path,
                    "original_filename": original_filename,
                    "stored_filename": stored_filename,
                    "file_extension": suffix.lstrip('.'),
                    "file_size_bytes": file_size_bytes,
                    "sheet_name": sheet_name,
                    "worksheet_index": index,
                    "worksheet_count": sheet_count,
                    "workbook_group_id": workbook_group_id,
                    "description_meta": build_description_meta(DESCRIPTION_SOURCE_DEFAULT),
                })
                payloads.append({
                    "datasource_name": datasource_name,
                    "datasource_schema": datasource_schema,
                    "datasource_config_json": datasource_config_json,
                    "knowledge_tag": f"upload_{workbook_group_id}_{index}",
                    "sheet_name": sheet_name,
                    "preview_summary": preview_summary,
                    "base_description": base_description,
                    "schema_hash": schema_guard_hash(safe_json_loads(datasource_schema, {})),
                })
            return payloads

        datasource_name = self._build_unique_datasource_name(
            insight_namespace_id=insight_namespace_id,
            base_name=base_name,
        )
        datasource_schema, preview_summary, base_description = self._build_file_datasource_schema_with_preview(
            stored_path,
            datasource_name,
        )
        datasource_config_json = dump_json({
            "file_path": file_path,
            "original_filename": original_filename,
            "stored_filename": stored_filename,
            "file_extension": suffix.lstrip('.'),
            "file_size_bytes": file_size_bytes,
            "description_meta": build_description_meta(DESCRIPTION_SOURCE_DEFAULT),
        })
        return [{
            "datasource_name": datasource_name,
            "datasource_schema": datasource_schema,
            "datasource_config_json": datasource_config_json,
            "knowledge_tag": f"upload_{uuid4().hex[:16]}",
            "sheet_name": "",
            "preview_summary": preview_summary,
            "base_description": base_description,
            "schema_hash": schema_guard_hash(safe_json_loads(datasource_schema, {})),
        }]

    def _list_excel_sheet_names(self, file_path: Path) -> list[str]:
        try:
            if file_path.suffix.lower() == '.xlsx':
                with pd.ExcelFile(file_path, engine='openpyxl') as excel_file:
                    return [str(name).strip() for name in excel_file.sheet_names if str(name).strip()]
            if file_path.suffix.lower() == '.xls':
                with pd.ExcelFile(file_path, engine='xlrd') as excel_file:
                    return [str(name).strip() for name in excel_file.sheet_names if str(name).strip()]
            raise ValueError(f'暂不支持解析 {file_path.suffix} 文件')
        except ImportError as exc:
            raise ValueError(f'文件解析依赖缺失：{exc}') from exc
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError(f'无法识别 Excel 工作表信息：{exc}') from exc

    def _build_upload_response_data(self, created_datasources: list[dict[str, Any]]) -> dict[str, Any]:
        if not created_datasources:
            return {}
        payload = dict(created_datasources[0])
        payload["created_count"] = len(created_datasources)
        payload["created_datasources"] = created_datasources
        return payload

    def _read_file_preview(self, file_path: Path, sheet_name: str | None = None) -> pd.DataFrame:
        """按文件类型读取预览数据；Excel 支持指定 sheet。"""
        try:
            if file_path.suffix.lower() == '.csv':
                return self._read_csv_preview(file_path)
            if file_path.suffix.lower() == '.xlsx':
                return pd.read_excel(
                    file_path,
                    nrows=50,
                    engine='openpyxl',
                    sheet_name=sheet_name if sheet_name else 0,
                )
            if file_path.suffix.lower() == '.xls':
                return pd.read_excel(
                    file_path,
                    nrows=50,
                    engine='xlrd',
                    sheet_name=sheet_name if sheet_name else 0,
                )
            raise ValueError(f'暂不支持解析 {file_path.suffix} 文件')
        except ImportError as exc:
            raise ValueError(f'文件解析依赖缺失：{exc}') from exc
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError(f'文件解析失败：{exc}') from exc

    def _read_file_content_page(
        self,
        file_path: Path,
        sheet_name: str | None,
        page_no: int,
        page_size: int,
    ) -> dict[str, Any]:
        suffix = file_path.suffix.lower()
        offset = (page_no - 1) * page_size
        if suffix == '.csv':
            dataframe = self._read_csv_page(file_path, offset=offset, limit=page_size)
            total_rows = self._count_csv_data_rows(file_path)
        elif suffix in {'.xlsx', '.xls'}:
            dataframe = self._read_excel_page(file_path, sheet_name=sheet_name, offset=offset, limit=page_size)
            total_rows = self._count_excel_data_rows(file_path, sheet_name=sheet_name)
        else:
            raise ValueError(f'暂不支持读取 {suffix} 文件内容')

        total_pages = (total_rows + page_size - 1) // page_size if total_rows >= 0 else 0
        return {
            "columns": [str(column) for column in dataframe.columns],
            "rows": self._dataframe_to_content_rows(dataframe),
            "total_rows": total_rows,
            "total_pages": total_pages,
            "has_next": page_no < total_pages if total_pages > 0 else False,
            "has_prev": page_no > 1,
        }

    def _read_csv_page(self, file_path: Path, offset: int, limit: int) -> pd.DataFrame:
        last_error: Exception | None = None
        for encoding in ('utf-8', 'utf-8-sig', 'gb18030', 'gbk'):
            try:
                return pd.read_csv(
                    file_path,
                    nrows=limit,
                    skiprows=range(1, offset + 1) if offset > 0 else None,
                    encoding=encoding,
                )
            except UnicodeDecodeError as exc:
                last_error = exc
                continue
            except Exception as exc:
                raise ValueError(f'CSV 文件读取失败：{exc}') from exc
        raise ValueError(f'CSV 文件编码识别失败：{last_error}')

    def _read_excel_page(
        self,
        file_path: Path,
        sheet_name: str | None,
        offset: int,
        limit: int,
    ) -> pd.DataFrame:
        try:
            if file_path.suffix.lower() == '.xlsx':
                return pd.read_excel(
                    file_path,
                    nrows=limit,
                    skiprows=range(1, offset + 1) if offset > 0 else None,
                    engine='openpyxl',
                    sheet_name=sheet_name if sheet_name else 0,
                )
            if file_path.suffix.lower() == '.xls':
                return pd.read_excel(
                    file_path,
                    nrows=limit,
                    skiprows=range(1, offset + 1) if offset > 0 else None,
                    engine='xlrd',
                    sheet_name=sheet_name if sheet_name else 0,
                )
            raise ValueError(f'暂不支持读取 {file_path.suffix} 文件内容')
        except ImportError as exc:
            raise ValueError(f'文件解析依赖缺失：{exc}') from exc
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError(f'Excel 文件读取失败：{exc}') from exc

    def _count_csv_data_rows(self, file_path: Path) -> int:
        last_error: Exception | None = None
        for encoding in ('utf-8', 'utf-8-sig', 'gb18030', 'gbk'):
            try:
                with file_path.open('r', encoding=encoding, newline='') as file:
                    line_count = sum(1 for _ in file)
                return max(line_count - 1, 0)
            except UnicodeDecodeError as exc:
                last_error = exc
                continue
            except Exception as exc:
                raise ValueError(f'CSV 文件行数统计失败：{exc}') from exc
        raise ValueError(f'CSV 文件编码识别失败：{last_error}')

    def _count_excel_data_rows(self, file_path: Path, sheet_name: str | None = None) -> int:
        try:
            if file_path.suffix.lower() == '.xlsx':
                from openpyxl import load_workbook

                workbook = load_workbook(file_path, read_only=True, data_only=True)
                try:
                    worksheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
                    return max(int(worksheet.max_row or 0) - 1, 0)
                finally:
                    workbook.close()
            if file_path.suffix.lower() == '.xls':
                import xlrd

                workbook = xlrd.open_workbook(str(file_path), on_demand=True)
                try:
                    worksheet = workbook.sheet_by_name(sheet_name) if sheet_name else workbook.sheet_by_index(0)
                    return max(int(worksheet.nrows or 0) - 1, 0)
                finally:
                    workbook.release_resources()
            raise ValueError(f'暂不支持读取 {file_path.suffix} 文件内容')
        except ImportError as exc:
            raise ValueError(f'文件解析依赖缺失：{exc}') from exc
        except Exception as exc:
            raise ValueError(f'Excel 文件行数统计失败：{exc}') from exc

    def _dataframe_to_content_rows(self, dataframe: pd.DataFrame) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = []
        for row in dataframe.to_dict(orient='records'):
            rows.append({
                str(column): self._preview_value(value)
                for column, value in row.items()
            })
        return rows

    def _to_dict(self, rel: InsightNsRelDatasource, datasource: InsightDatasource) -> dict[str, Any]:
        datasource_config = safe_json_loads(datasource.datasource_config_json, {})
        return {
            "id": rel.id,
            "insight_namespace_id": rel.insight_namespace_id,
            "insight_conversation_id": rel.insight_conversation_id,
            "datasource_id": datasource.id,
            "datasource_type": datasource.datasource_type,
            "datasource_name": datasource.datasource_name,
            "knowledge_tag": datasource.knowledge_tag,
            "uns_node_id": datasource.uns_node_id,
            "datasource_schema": datasource.datasource_schema,
            "datasource_config_json": datasource.datasource_config_json,
            "sheet_name": str(datasource_config.get("sheet_name") or ''),
            "sort_no": rel.sort_no,
            "bind_source": rel.bind_source,
            "created_at": rel.created_at.isoformat() if rel.created_at else None,
            "updated_at": rel.updated_at.isoformat() if rel.updated_at else None,
        }

    def _datasource_to_dict(
        self,
        datasource: InsightDatasource,
        checked: bool = False,
        insight_conversation_id: int = 0,
    ) -> dict[str, Any]:
        datasource_config = safe_json_loads(datasource.datasource_config_json, {})
        return {
            "id": datasource.id,
            "datasource_id": datasource.id,
            "insight_namespace_id": datasource.insight_namespace_id,
            "insight_conversation_id": insight_conversation_id,
            "datasource_type": datasource.datasource_type,
            "datasource_name": datasource.datasource_name,
            "knowledge_tag": datasource.knowledge_tag,
            "uns_node_id": datasource.uns_node_id,
            "datasource_schema": datasource.datasource_schema,
            "datasource_config_json": datasource.datasource_config_json,
            "sheet_name": str(datasource_config.get("sheet_name") or ''),
            "checked": checked,
            "created_at": datasource.created_at.isoformat() if datasource.created_at else None,
            "updated_at": datasource.updated_at.isoformat() if datasource.updated_at else None,
        }
